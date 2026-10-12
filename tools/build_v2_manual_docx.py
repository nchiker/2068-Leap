#!/usr/bin/env python3
"""Build the Word edition of the V2 user manual (docs/v2/manual/*.md -> docs/v2/2068_Leap_V2_Users_Manual.docx).

Adapted from tools/build_user_manual_docx.py (the V1 manual's builder, whose helpers it reuses): the seven
Markdown parts are concatenated in order, headings become Word headings with bookmarks (so the in-manual
links work), tables, lists, code/output blocks and the screen pictures are carried over.

    python3 tools/build_v2_manual_docx.py [output.docx]      (make manual-v2)
"""
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_user_manual_docx as B  # noqa: E402  (helpers: configure, shade, bookmarks, fields)
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MANUAL = ROOT / "docs" / "v2" / "manual"
OUTPUT = ROOT / "docs" / "v2" / "2068_Leap_V2_Users_Manual.docx"
PARTS = ["01_introduction.md", "02_beginners_guide.md", "03_keywords.md", "04_concepts.md",
         "05_packages.md", "06_examples.md", "07_appendices.md"]
FLOW_PARTS = {"03_keywords.md", "04_concepts.md"}      # entries flow on: no page break before each ##

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<![\w*])\*[^*\n]+\*(?![\w*])|\[[^]]+\]\([^)]+\))")


def slug(text):
    return B.markdown_anchor(text)


def heading_anchors():
    out = set()
    for name in PARTS:
        in_code = False
        for line in (MANUAL / name).read_text(encoding="utf-8").splitlines():
            if line.startswith("```"):
                in_code = not in_code
            m = None if in_code else re.match(r"^#{1,3}\s+(.+)$", line)
            if m:
                out.add(slug(m.group(1)))
                first = re.match(r"^([A-Z$ ]+?)\s{2}\[package", m.group(1))
                if first:
                    out.add(slug(first.group(1)))
    return out


ANCHORS = heading_anchors()


def add_inline(paragraph, text):
    cursor = 0
    for match in INLINE.finditer(text):
        paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("`"):
            paragraph.add_run(token[1:-1]).style = "Code Inline"
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            first = len(paragraph.runs)
            add_inline(paragraph, token[1:-1])
            for run in paragraph.runs[first:]:
                run.italic = True
        else:
            label, target = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            anchor = target.partition("#")[2]
            if anchor and anchor in ANCHORS:
                B.add_internal_hyperlink(paragraph, label, "#" + anchor)
            else:
                paragraph.add_run(label)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def parse(document, lines, flow):
    index, in_code, code_lines, info = 0, False, [], ""
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                B.shade(paragraph, "E8F1E4" if info.startswith(("output", "screen")) else "F2F4F7")
                code_lines, in_code = [], False
            else:
                in_code, info = True, line[3:].strip()
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line or line == "---":
            index += 1
            continue
        image = re.match(r"^!\[([^]]*)\]\(([^)]+)\)\s*$", line)
        if image:
            path = MANUAL / image.group(2)
            if path.exists():
                document.add_picture(str(path), width=Inches(4.0))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(image.group(1))
                run.italic = True
                run.font.size = Pt(9)
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_heading(level=level)
            if level == 1 or (level == 2 and not flow):
                paragraph.paragraph_format.page_break_before = True
            add_inline(paragraph, heading.group(2))
            B.add_bookmark(paragraph, slug(heading.group(2)))
            first = re.match(r"^([A-Z$ ]+?)\s{2}\[package", heading.group(2))
            if first:
                B.add_bookmark(paragraph, slug(first.group(1)))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([p.strip() for p in re.split(r"(?<!\\)\|", lines[index].strip().strip("|"))])
                index += 1
            rows.pop(1)
            width = max(map(len, rows))
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Light Shading Accent 1"
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    add_inline(cell.paragraphs[0], value.replace("\\|", "|"))
                    if r == 0:
                        B.shade(cell, "1F4E79")
                        for run in cell.paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
            document.add_paragraph()
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or number:
            item = (bullet or number).group(1)
            index += 1
            while index < len(lines) and re.match(r"^\s{2,}\S", lines[index]):
                item += " " + lines[index].strip()
                index += 1
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, item)
            continue
        parts = [line.strip()]
        index += 1
        while index < len(lines):
            cand = lines[index]
            if (not cand or cand.startswith(("#", "```", "|", "- ", "* ", "![")) or cand == "---"
                    or re.match(r"^\d+\.\s", cand)):
                break
            parts.append(cand.strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(parts))


def main():
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    document = Document()
    B.configure(document)
    for sec in document.sections:
        footer = sec.footer.paragraphs[0]
        for run in list(footer.runs):
            run._r.getparent().remove(run._r)
        footer.add_run("2068-Leap V2 User's Manual   •   ")
        B.add_field(footer, "PAGE")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2068-Leap V2")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("User Manual and Language Reference\nTimex Sinclair 2068")
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    document.add_paragraph("\n")
    ed = document.add_paragraph()
    ed.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ed.add_run("V2 Enhanced -- Emulator-Qualified Release (v2.0.0-emu.1) • %s" % date.today().isoformat()).italic = True
    document.add_page_break()
    document.add_heading("Contents", level=1)
    toc = document.add_paragraph()
    B.add_field(toc, 'TOC \\o "1-2" \\h \\z \\u')
    note = document.add_paragraph("If page numbers are not visible, open the document in Word or LibreOffice "
                                  "and update the table of contents field.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)

    for name in PARTS:
        text = (MANUAL / name).read_text(encoding="utf-8").splitlines()
        parse(document, text, flow=name in FLOW_PARTS)
    document.core_properties.title = "2068-Leap V2 User Manual and Language Reference"
    document.core_properties.subject = "2068-Leap V2 Enhanced, Emulator-Qualified Release"
    document.core_properties.author = "2068-Leap Project"
    document.core_properties.keywords = "TS2068, Timex Sinclair, BASIC, ROM, V2"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print("Wrote %s" % output)


if __name__ == "__main__":
    main()
